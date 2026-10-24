import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent

# Color Palette
DARK_BG = RGBColor(10, 31, 26)       # Deep Forest Dark #0a1f1a
CARD_BG = RGBColor(18, 48, 40)       # Dark Card
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(167, 243, 208)  # Soft mint green
EMERALD = RGBColor(16, 185, 129)     # Brand Emerald #10b981
EMERALD_LIGHT = RGBColor(209, 250, 229)
LIGHT_BG = RGBColor(248, 250, 252)   # Slate 50
TEXT_DARK = RGBColor(15, 23, 42)     # Slate 900
TEXT_GRAY = RGBColor(100, 116, 139)  # Slate 500
BORDER_COLOR = RGBColor(226, 232, 240)
ACCENT_CARD = RGBColor(236, 253, 245)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_text="PLANT DISEASE DETECTION AI"):
    # Header badge / category
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11), Inches(0.35))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = EMERALD
    p_c.font.name = "Segoe UI"

    # Main Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.8))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_DARK
    p_t.font.name = "Segoe UI"

def set_slide_bg(slide, rgb_color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = rgb_color

# ==========================================
# SLIDE 1: Title Slide (Dark Theme)
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_bg(s1, DARK_BG)

# Title container shape
shape1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.733), Inches(5.1))
shape1.fill.solid()
shape1.fill.fore_color.rgb = CARD_BG
shape1.line.color.rgb = RGBColor(30, 80, 65)
shape1.line.width = Pt(1.5)

tb1 = s1.shapes.add_textbox(Inches(1.3), Inches(1.6), Inches(10.7), Inches(4.3))
tf1 = tb1.text_frame
tf1.word_wrap = True

p1 = tf1.paragraphs[0]
p1.text = "🌿 FINAL YEAR CAPSTONE PROJECT"
p1.font.size = Pt(12)
p1.font.bold = True
p1.font.color.rgb = EMERALD
p1.font.name = "Segoe UI"
p1.space_after = Pt(14)

p2 = tf1.add_paragraph()
p2.text = "Plant Disease Detection System"
p2.font.size = Pt(36)
p2.font.bold = True
p2.font.color.rgb = TEXT_WHITE
p2.font.name = "Segoe UI"
p2.space_after = Pt(8)

p3 = tf1.add_paragraph()
p3.text = "Autonomous Tomato Leaf Disease Classification using Deep Learning (CNN) & Web Deployment"
p3.font.size = Pt(17)
p3.font.color.rgb = TEXT_MUTED
p3.font.name = "Segoe UI"
p3.space_after = Pt(36)

p4 = tf1.add_paragraph()
p4.text = "Model: 89.10% Validation Accuracy  |  10 Foliar Classes  |  FastAPI & Tailwind Web App"
p4.font.size = Pt(13)
p4.font.bold = True
p4.font.color.rgb = EMERALD
p4.font.name = "Segoe UI"

# ==========================================
# SLIDE 2: Problem Statement & Motivation
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_bg(s2, LIGHT_BG)
add_header(s2, "Problem Statement & Project Motivation", "BACKGROUND & CONTEXT")

cards_data_2 = [
    ("Agricultural Impact", "Tomato is one of the most cultivated crops globally, but foliar diseases like Early Blight, Late Blight, and Mosaic Virus cause up to 30%-40% yield losses every season."),
    ("Manual Inspection Limits", "Traditional detection depends heavily on visual inspection by agricultural experts, which is slow, error-prone, subjective, and unavailable in remote rural regions."),
    ("Need for Precision AI", "An automated, camera-based AI diagnostic tool empowers farmers to detect infections early, stop pathogen spread, and apply precise treatments without excessive chemical waste.")
]

for idx, (head, body) in enumerate(cards_data_2):
    x = Inches(0.8 + idx * 3.95)
    c_shape = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.75), Inches(4.8))
    c_shape.fill.solid()
    c_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    c_shape.line.color.rgb = BORDER_COLOR
    c_shape.line.width = Pt(1)

    tb = s2.shapes.add_textbox(x + Inches(0.3), Inches(2.1), Inches(3.15), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = f"0{idx+1}"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(10)

    p_h = tf.add_paragraph()
    p_h.text = head
    p_h.font.size = Pt(16)
    p_h.font.bold = True
    p_h.font.color.rgb = TEXT_DARK
    p_h.space_after = Pt(12)

    p_b = tf.add_paragraph()
    p_b.text = body
    p_b.font.size = Pt(12)
    p_b.font.color.rgb = TEXT_GRAY
    p_b.line_spacing = 1.3

# ==========================================
# SLIDE 3: Project Objectives
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_bg(s3, LIGHT_BG)
add_header(s3, "Key Project Objectives & Deliverables", "GOALS & SCOPE")

objs = [
    ("🎯 Deep Learning Classification", "Design and train an optimized Convolutional Neural Network (CNN) in TensorFlow/Keras to accurately classify tomato leaves into 10 disease & healthy categories."),
    ("📊 High Validation Benchmark", "Achieve competitive diagnostic accuracy (>85%) evaluated on a distinct, unseen validation dataset of 1,000 images."),
    ("💻 Full-Stack Web Application", "Engineer an interactive, responsive web portal (FastAPI + Tailwind CSS) supporting drag-and-drop file upload and live webcam capture."),
    ("💊 Agronomic Treatment Advisory", "Integrate automated biological, organic, and chemical remedy recommendations alongside pathogen etiology for every detected disease."),
    ("🚀 Instant Cloud Deployment", "Containerize via Docker and deploy publicly on Hugging Face Spaces & Cloudflare Tunnel for worldwide accessibility on any mobile device.")
]

for i, (head, text) in enumerate(objs):
    y = Inches(1.7 + i * 1.05)
    bar = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(0.92))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bar.line.color.rgb = BORDER_COLOR

    tb = s3.shapes.add_textbox(Inches(1.0), y + Inches(0.08), Inches(11.3), Inches(0.76))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = head
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    
    p_sub = tf.add_paragraph()
    p_sub.text = text
    p_sub.font.size = Pt(10.5)
    p_sub.font.color.rgb = TEXT_GRAY

# ==========================================
# SLIDE 4: Dataset & Classes Overview
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_bg(s4, LIGHT_BG)
add_header(s4, "Dataset & Supported Disease Categories", "DATASET SPECIFICATIONS")

# Left box: 10 classes
c_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(6.8), Inches(4.9))
c_box.fill.solid()
c_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
c_box.line.color.rgb = BORDER_COLOR

tb_c = s4.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(6.2), Inches(4.5))
tf_c = tb_c.text_frame
tf_c.word_wrap = True

p_ch = tf_c.paragraphs[0]
p_ch.text = "10 Targeted Tomato Leaf Conditions:"
p_ch.font.size = Pt(14)
p_ch.font.bold = True
p_ch.font.color.rgb = TEXT_DARK
p_ch.space_after = Pt(10)

classes_list = [
    "1. Bacterial Spot (Xanthomonas spp.)",
    "2. Early Blight (Alternaria solani)",
    "3. Late Blight (Phytophthora infestans)",
    "4. Leaf Mold (Passalora fulva)",
    "5. Septoria Leaf Spot (Septoria lycopersici)",
    "6. Two-Spotted Spider Mite (Tetranychus urticae)",
    "7. Target Spot (Corynespora cassiicola)",
    "8. Tomato Yellow Leaf Curl Virus (TYLCV)",
    "9. Tomato Mosaic Virus (ToMV)",
    "10. Healthy Tomato Foliage"
]

for c in classes_list:
    p_li = tf_c.add_paragraph()
    p_li.text = c
    p_li.font.size = Pt(10.5)
    p_li.font.color.rgb = TEXT_DARK if "Healthy" not in c else EMERALD
    p_li.font.bold = "Healthy" in c

# Right box: Dataset specs & sample images
r_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.8), Inches(4.733), Inches(4.9))
r_box.fill.solid()
r_box.fill.fore_color.rgb = ACCENT_CARD
r_box.line.color.rgb = EMERALD
r_box.line.width = Pt(1)

tb_r = s4.shapes.add_textbox(Inches(8.0), Inches(2.0), Inches(4.3), Inches(4.5))
tf_r = tb_r.text_frame
tf_r.word_wrap = True

p_rh = tf_r.paragraphs[0]
p_rh.text = "Dataset Characteristics:"
p_rh.font.size = Pt(14)
p_rh.font.bold = True
p_rh.font.color.rgb = TEXT_DARK
p_rh.space_after = Pt(8)

specs = [
    "• Input Resolution: 128 × 128 × 3 RGB channels",
    "• Data Preprocessing: Rescaling, normalization",
    "• Augmentation: Random rotations, flips, zoom",
    "• Validation Split: 1,000 distinct holdout images",
    "• Balance: 100 images per class in validation set",
    "• High quality foliar photographic samples"
]

for s in specs:
    p_s = tf_r.add_paragraph()
    p_s.text = s
    p_s.font.size = Pt(11)
    p_s.font.color.rgb = TEXT_GRAY
    p_s.space_after = Pt(4)

# Insert sample images if present
sample_early = ROOT / "static" / "samples" / "Tomato___Early_blight.jpg"
sample_healthy = ROOT / "static" / "samples" / "Tomato___healthy.jpg"
if sample_early.exists() and sample_healthy.exists():
    s4.shapes.add_picture(str(sample_early), Inches(8.1), Inches(4.6), width=Inches(1.8), height=Inches(1.8))
    s4.shapes.add_picture(str(sample_healthy), Inches(10.2), Inches(4.6), width=Inches(1.8), height=Inches(1.8))

# ==========================================
# SLIDE 5: CNN Architecture & Pipeline
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_bg(s5, LIGHT_BG)
add_header(s5, "Deep CNN Model Architecture & Pipeline", "NEURAL NETWORK DESIGN")

steps = [
    ("1. Image Input", "128×128×3 RGB foliar image uploaded via web or camera."),
    ("2. Feature Extraction", "Sequential Conv2D layers with ReLU activations extract edge, texture, and lesion spots."),
    ("3. Spatial Pooling", "MaxPooling2D layers downsample spatial dimensions while retaining invariant patterns."),
    ("4. Regularization", "Dropout layers mitigate overfitting and improve real-world generalization."),
    ("5. Dense & Softmax", "Flatten layer followed by Dense layers; Softmax outputs 10 probability scores.")
]

for idx, (st_t, st_d) in enumerate(steps):
    x = Inches(0.8 + idx * 2.38)
    shape_s = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(2.25), Inches(4.5))
    shape_s.fill.solid()
    shape_s.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape_s.line.color.rgb = EMERALD if idx == 4 else BORDER_COLOR
    shape_s.line.width = Pt(1.5) if idx == 4 else Pt(1)

    tb = s5.shapes.add_textbox(x + Inches(0.15), Inches(2.2), Inches(1.95), Inches(4.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p_num = tf.paragraphs[0]
    p_num.text = f"LAYER 0{idx+1}"
    p_num.font.size = Pt(10)
    p_num.font.bold = True
    p_num.font.color.rgb = EMERALD
    p_num.space_after = Pt(8)

    p_h = tf.add_paragraph()
    p_h.text = st_t
    p_h.font.size = Pt(13)
    p_h.font.bold = True
    p_h.font.color.rgb = TEXT_DARK
    p_h.space_after = Pt(10)

    p_b = tf.add_paragraph()
    p_b.text = st_d
    p_b.font.size = Pt(10.5)
    p_b.font.color.rgb = TEXT_GRAY

# ==========================================
# SLIDE 6: Model Evaluation & Results
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_bg(s6, LIGHT_BG)
add_header(s6, "Experimental Evaluation & Performance Metrics", "MODEL VALIDATION")

# Left Column: Key Metric Numbers
metrics_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9))
metrics_box.fill.solid()
metrics_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
metrics_box.line.color.rgb = BORDER_COLOR

tb_m = s6.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
tf_m = tb_m.text_frame
tf_m.word_wrap = True

p_mh = tf_m.paragraphs[0]
p_mh.text = "Validation Set Benchmark"
p_mh.font.size = Pt(15)
p_mh.font.bold = True
p_mh.font.color.rgb = TEXT_DARK
p_mh.space_after = Pt(14)

m_stats = [
    ("Validation Accuracy", "89.10%", "High precision on 1,000 unseen images"),
    ("Macro F1-Score", "0.89", "Balanced across all 10 leaf classes"),
    ("Inference Latency", "< 300 ms", "Optimized for instant interactive web diagnosis"),
    ("Top Performing Class", "Mosaic Virus (99%)", "Near-perfect virus identification")
]

for label, val, note in m_stats:
    p_l = tf_m.add_paragraph()
    p_l.text = f"{label}: {val}"
    p_l.font.size = Pt(12)
    p_l.font.bold = True
    p_l.font.color.rgb = EMERALD if "%" in val else TEXT_DARK
    
    p_n = tf_m.add_paragraph()
    p_n.text = f"  └ {note}"
    p_n.font.size = Pt(9.5)
    p_n.font.color.rgb = TEXT_GRAY
    p_n.space_after = Pt(8)

# Right Column: Embed Confusion Matrix Image
cm_path = ROOT / "models" / "confusion_matrix.png"
if cm_path.exists():
    s6.shapes.add_picture(str(cm_path), Inches(6.8), Inches(1.8), width=Inches(5.7), height=Inches(4.9))

# ==========================================
# SLIDE 7: Web Application Features
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_bg(s7, LIGHT_BG)
add_header(s7, "Full-Stack Web Application Features", "USER INTERFACE & EXPERIENCE")

feat_data = [
    ("📸 Drag & Drop + Webcam", "Users can drag and drop high-res leaf photos or capture leaves directly using mobile/laptop webcams with real-time framing."),
    ("⚡ Instant Inference Engine", "FastAPI asynchronous backend delivers sub-second predictions and generates animated scanning visuals."),
    ("📊 Class Probability Radar", "Interactive Chart.js visualizations display confidence percentages across all 10 possible conditions for transparency."),
    ("🧪 Agronomic Treatment Plans", "Actionable advice split into Symptoms, Eco-Friendly Organic Remedies, Commercial Chemical Sprays, and Cultural Prevention."),
    ("📚 Disease Encyclopedia", "Searchable digital library covering pathogen etiology, vectors, severity ratings, and management for all tomato diseases."),
    ("🖨️ Formatted Pathology Reports", "One-click generation of printable diagnostic health certificates for farm records or academic evaluation.")
]

for idx, (f_title, f_desc) in enumerate(feat_data):
    col = idx % 3
    row = idx // 3
    x = Inches(0.8 + col * 3.95)
    y = Inches(1.8 + row * 2.5)

    f_shape = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.75), Inches(2.25))
    f_shape.fill.solid()
    f_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    f_shape.line.color.rgb = BORDER_COLOR

    tb = s7.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), Inches(3.35), Inches(1.95))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = f_title
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(6)

    p_d = tf.add_paragraph()
    p_d.text = f_desc
    p_d.font.size = Pt(10)
    p_d.font.color.rgb = TEXT_GRAY
    p_d.line_spacing = 1.25

# ==========================================
# SLIDE 8: Technology Stack
# ==========================================
s8 = prs.slides.add_slide(blank_layout)
set_slide_bg(s8, LIGHT_BG)
add_header(s8, "Technologies & Tools Used", "SYSTEM IMPLEMENTATION")

techs = [
    ("Machine Learning & AI", "Python 3.12, TensorFlow 2.21, Keras 3, NumPy, Pillow, Scikit-learn, Matplotlib"),
    ("Backend API Server", "FastAPI, Uvicorn (ASGI), Python-Multipart, Starlette"),
    ("Frontend Interface", "HTML5, Tailwind CSS, Chart.js, Lucide Icons, Vanilla JavaScript (SPA)"),
    ("Deployment & DevOps", "Docker containerization, Hugging Face Spaces (16GB RAM), Cloudflare Tunnel (HTTPS)"),
    ("Development & Versioning", "Git, GitHub, Visual Studio Code, Python Virtual Environments (.venv)")
]

for idx, (layer, tools) in enumerate(techs):
    y = Inches(1.8 + idx * 1.0)
    b_shape = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(0.85))
    b_shape.fill.solid()
    b_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    b_shape.line.color.rgb = BORDER_COLOR

    tb = s8.shapes.add_textbox(Inches(1.1), y + Inches(0.08), Inches(11.1), Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = layer
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    p_t = tf.add_paragraph()
    p_t.text = tools
    p_t.font.size = Pt(10.5)
    p_t.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 9: Limitations & Future Scope
# ==========================================
s9 = prs.slides.add_slide(blank_layout)
set_slide_bg(s9, LIGHT_BG)
add_header(s9, "Current Limitations & Future Enhancements", "PROJECT ROADMAP")

# Left: Limitations
l_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9))
l_box.fill.solid()
l_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
l_box.line.color.rgb = BORDER_COLOR

tb_l = s9.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
tf_l = tb_l.text_frame
tf_l.word_wrap = True

p_lh = tf_l.paragraphs[0]
p_lh.text = "Current Limitations:"
p_lh.font.size = Pt(14)
p_lh.font.bold = True
p_lh.font.color.rgb = RGBColor(225, 29, 72) # Rose
p_lh.space_after = Pt(10)

lims = [
    "• Single-Crop Focus: Trained exclusively on tomato foliage.",
    "• Environmental Sensitivity: Harsh shadows, extreme glare, or blurred images may reduce prediction confidence.",
    "• Co-infections: Difficult to separate multiple simultaneous pathogens on a single leaflet.",
    "• Laboratory Diagnosis: AI serves as decision support, not certified chemical laboratory verification."
]
for l in lims:
    p = tf_l.add_paragraph()
    p.text = l
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_GRAY
    p.space_after = Pt(6)

# Right: Future Scope
r_box2 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.633), Inches(4.9))
r_box2.fill.solid()
r_box2.fill.fore_color.rgb = ACCENT_CARD
r_box2.line.color.rgb = EMERALD

tb_r2 = s9.shapes.add_textbox(Inches(7.2), Inches(2.0), Inches(5.0), Inches(4.5))
tf_r2 = tb_r2.text_frame
tf_r2.word_wrap = True

p_rh2 = tf_r2.paragraphs[0]
p_rh2.text = "Future Enhancements:"
p_rh2.font.size = Pt(14)
p_rh2.font.bold = True
p_rh2.font.color.rgb = EMERALD
p_rh2.space_after = Pt(10)

futs = [
    "• Multi-Crop Expansion: Extend dataset to potato, pepper, corn, apple, and grape leaves.",
    "• Edge AI & Mobile Apps: Quantize model with TensorFlow Lite for offline on-device scanning in remote fields.",
    "• IoT Drone Integration: Aerial crop canopy scanning using automated UAVs.",
    "• Weather-Driven Risk Warning: Combine humidity/temperature forecasts with disease likelihood models."
]
for f in futs:
    p = tf_r2.add_paragraph()
    p.text = f
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(6)

# ==========================================
# SLIDE 10: Conclusion & Summary
# ==========================================
s10 = prs.slides.add_slide(blank_layout)
set_slide_bg(s10, DARK_BG)

c_shape10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.733), Inches(5.1))
c_shape10.fill.solid()
c_shape10.fill.fore_color.rgb = CARD_BG
c_shape10.line.color.rgb = RGBColor(30, 80, 65)
c_shape10.line.width = Pt(1.5)

tb10 = s10.shapes.add_textbox(Inches(1.3), Inches(1.6), Inches(10.7), Inches(4.3))
tf10 = tb10.text_frame
tf10.word_wrap = True

p10_1 = tf10.paragraphs[0]
p10_1.text = "CONCLUSION & SUMMARY"
p10_1.font.size = Pt(12)
p10_1.font.bold = True
p10_1.font.color.rgb = EMERALD
p10_1.space_after = Pt(12)

p10_2 = tf10.add_paragraph()
p10_2.text = "Bridging Deep Learning & Precision Agriculture"
p10_2.font.size = Pt(30)
p10_2.font.bold = True
p10_2.font.color.rgb = TEXT_WHITE
p10_2.space_after = Pt(16)

p10_3 = tf10.add_paragraph()
p10_3.text = "• Successfully trained a CNN achieving 89.10% validation accuracy across 10 tomato leaf classes.\n• Built an intuitive, full-featured web portal with live camera diagnostics and agronomic treatment plans.\n• Provides instant actionable guidance for farmers, minimizing yield loss and chemical overuse.\n• Ready for production deployment on Docker and Hugging Face Spaces."
p10_3.font.size = Pt(13)
p10_3.font.color.rgb = TEXT_MUTED
p10_3.line_spacing = 1.3
p10_3.space_after = Pt(24)

p10_4 = tf10.add_paragraph()
p10_4.text = "Thank You! Questions & Discussion Welcome."
p10_4.font.size = Pt(15)
p10_4.font.bold = True
p10_4.font.color.rgb = EMERALD

output_path = ROOT / "Plant_Disease_Detection_Presentation.pptx"
prs.save(str(output_path))
print(f"Presentation saved successfully at: {output_path}")
